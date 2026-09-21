"""Read an XLSX whitepaper without modifying it; import merged values explicitly.

Usage: python3 assets/metric-center/import_whitepaper.py /path/to/whitepaper.xlsx
Requires openpyxl only when importing; normal site builds use the committed JSON.
"""
from pathlib import Path
import argparse
import hashlib
import json
import openpyxl

ROOT = Path(__file__).resolve().parent


def records(sheet, keys, required):
    merged = {}
    for area in sheet.merged_cells.ranges:
        for row in sheet.iter_rows(min_row=area.min_row, max_row=area.max_row,
                                   min_col=area.min_col, max_col=area.max_col):
            for cell in row:
                merged[cell.coordinate] = sheet.cell(area.min_row, area.min_col)
    result = []
    for row in sheet.iter_rows(min_row=2, max_col=len(keys)):
        data, cells = {}, {}
        for key, cell in zip(keys, row):
            source = merged.get(cell.coordinate, cell)
            data[key] = str(source.value) if source.value is not None else ''
            cells[key] = source.coordinate
        if not data[required].strip():
            continue
        data.update(row=row[0].row, sheet=sheet.title, cells=cells)
        result.append(data)
    return result


def build(path):
    book = openpyxl.load_workbook(path, data_only=True)
    metrics = records(book['任务结束上报指标'],
                      ['category', 'name', 'definition', 'bi', 'dashboard', 'url',
                       'table', 'report', 'pc', 'mobile', 'sql'], 'name')
    dimensions = records(book['维度字典'],
                         ['name', 'field', 'definition', 'values', 'pc', 'mobile',
                          'scope', 'sql'], 'name')
    for metric in metrics:
        metric['id'] = 'wp-' + hashlib.sha256(metric['name'].encode()).hexdigest()[:12]
    result = dict(source=path.name, sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
                  metrics=metrics, dimensions=dimensions,
                  emptySheets=[s.title for s in book if not any(c.value is not None for row in s for c in row)])
    (ROOT / 'whitepaper.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
    print(f"Imported {len(metrics)} metrics and {len(dimensions)} dimensions")


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source', type=Path)
    build(parser.parse_args().source)
